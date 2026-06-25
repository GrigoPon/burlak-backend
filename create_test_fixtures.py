from openpyxl import Workbook
from pathlib import Path
from zipfile import ZipFile

# Создаём папку fixtures
fixtures_dir = Path('tests/fixtures')
fixtures_dir.mkdir(parents=True, exist_ok=True)

print('Папка fixtures готова')

# ==================== 1. minimal_bom.xlsx ====================
wb = Workbook()
ws = wb.active
ws.title = 'BOM'

headers = ['Артикул', 'Наименование', 'Количество', 'Цена', 'Поставщик', 'Ссылка']
for col, header in enumerate(headers, 1):
    ws.cell(row=1, column=col, value=header)

data = [
    ['ART-001', 'Товар 1', 5, 1500, 'Поставщик A', 'https://example.com/1'],
    ['ART-002', 'Товар 2', 3, 2500, 'Поставщик B', 'https://example.com/2'],
    ['ART-003', 'Товар 3', 10, 800, 'Поставщик A', ''],
]

for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

wb.save(fixtures_dir / 'minimal_bom.xlsx')
print('✅ minimal_bom.xlsx создан')

# ==================== 2. test_card.xlsx ====================
wb = Workbook()
ws = wb.active
ws.title = 'Card'

headers = ['Поле', 'Значение']
for col, header in enumerate(headers, 1):
    ws.cell(row=1, column=col, value=header)

data = [
    ['Артикул', 'TEST-CARD-001'],
    ['Название', 'Тестовая карточка товара'],
    ['Описание', 'Это тестовое описание для проверки парсера карточек'],
    ['Цена', 12990],
    ['Вес', 0.85],
    ['Категория', 'Электроника/Смартфоны'],
]

for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

wb.save(fixtures_dir / 'test_card.xlsx')
print('✅ test_card.xlsx создан')

# ==================== 3. test_archive.zip ====================
with ZipFile(fixtures_dir / 'test_archive.zip', 'w') as zipf:
    zipf.write(fixtures_dir / 'test_card.xlsx', 'cards/card_001.xlsx')
    zipf.write(fixtures_dir / 'test_card.xlsx', 'cards/card_002.xlsx')
    zipf.writestr('readme.txt', 'Тестовый архив для интеграционных тестов Burlak Parser')
    zipf.writestr('cards/card_003.xlsx', 'dummy content for testing')

print('✅ test_archive.zip создан')
print('Все тестовые файлы успешно созданы!')
