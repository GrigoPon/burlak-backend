"""Unit-тесты для splitter.py.

Покрытие:
  - CardSplitter.split_file: базовое разделение, один лист, несколько листов
  - CardSplitter.split_file: edge cases (не .xlsx, лист не найден, дубликаты имён)
  - CardSplitter._extract_sheet_via_zip: проверка содержимого split-файлов
  - CardSplitter.split_many_parallel: параллельное разделение
  - _clean_named_ranges: очистка named ranges
  - Интеграция: валидность .xlsx после разделения
"""

from __future__ import annotations

import os
import shutil
import tempfile
from typing import Any, Dict, List, Optional

import openpyxl
import pytest

from burlak_parser.splitter import (
    CardSplitter,
    _clean_named_ranges,
    _split_file_worker,
)
from burlak_parser.heuristic_analyzer import HeuristicAnalyzer


# ═══════════════════════════════════════════════════════════════════════
#  ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ═══════════════════════════════════════════════════════════════════════

def _create_multi_sheet_xlsx(
    dir_path: str,
    filename: str = "test_multi.xlsx",
    sheets: Optional[Dict[str, List[List[Any]]]] = None,
) -> str:
    """Create an .xlsx file with multiple sheets and return its path."""
    if sheets is None:
        sheets = {
            "Sheet1": [["A1", "B1"], ["A2", "B2"]],
            "Sheet2": [["C1", "D1"], ["C2", "D2"]],
            "Sheet3": [["E1", "F1"], ["E2", "F2"]],
        }
    path = os.path.join(dir_path, filename)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(title=name)
        for r_idx, row in enumerate(data, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(path)
    return path


def _count_xlsx_sheets(path: str) -> int:
    """Count sheets in an .xlsx file using openpyxl."""
    wb = openpyxl.load_workbook(path)
    count = len(wb.sheetnames)
    wb.close()
    return count


def _get_xlsx_cell(path: str, sheet: str, row: int, col: int) -> Any:
    """Get cell value from an .xlsx file."""
    wb = openpyxl.load_workbook(path)
    val = wb[sheet].cell(row=row, column=col).value
    wb.close()
    return val


# ═══════════════════════════════════════════════════════════════════════
#  Fixtures
# ═══════════════════════════════════════════════════════════════════════

@pytest.fixture
def tmp_dir() -> str:
    """Create a temporary directory for test files."""
    path = tempfile.mkdtemp(prefix="splitter_test_", dir="/home/rpaup/projects/burlak")
    yield path
    shutil.rmtree(path, ignore_errors=True)


@pytest.fixture
def multi_sheet_xlsx(tmp_dir: str) -> str:
    """Create a 3-sheet .xlsx file."""
    return _create_multi_sheet_xlsx(tmp_dir)


@pytest.fixture
def single_sheet_xlsx(tmp_dir: str) -> str:
    """Create a single-sheet .xlsx file."""
    return _create_multi_sheet_xlsx(
        tmp_dir, "single.xlsx",
        sheets={"OnlySheet": [["Data1"], ["Data2"]]},
    )


# ═══════════════════════════════════════════════════════════════════════
#  1. CardSplitter.split_file — basic
# ═══════════════════════════════════════════════════════════════════════

class TestSplitFileBasic:
    def test_split_one_sheet(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Split one sheet from a multi-sheet file."""
        output_dir = os.path.join(tmp_dir, "out1")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], "TestCard",
        )
        assert len(created) == 1, f"Expected 1 file, got {len(created)}"
        assert os.path.exists(created[0])
        # Verify the output file has only 1 sheet
        assert _count_xlsx_sheets(created[0]) == 1, "Split file should have 1 sheet"

    def test_split_multiple_sheets(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Split multiple sheets from a multi-sheet file."""
        output_dir = os.path.join(tmp_dir, "out2")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1", "Sheet2"], "TestCard",
        )
        assert len(created) == 2, f"Expected 2 files, got {len(created)}"
        for fp in created:
            assert os.path.exists(fp)
            assert _count_xlsx_sheets(fp) == 1, "Each split file should have 1 sheet"

    def test_split_all_sheets(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Split all 3 sheets."""
        output_dir = os.path.join(tmp_dir, "out3")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1", "Sheet2", "Sheet3"], "TestCard",
        )
        assert len(created) == 3, f"Expected 3 files, got {len(created)}"

    def test_data_preserved(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Data in split file should match original sheet data."""
        output_dir = os.path.join(tmp_dir, "out4")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet2"], "TestCard",
        )[0]

        # Sheet2 in original had C1, D1; C2, D2
        assert _count_xlsx_sheets(created) == 1
        split_wb = openpyxl.load_workbook(created)
        split_ws = split_wb.active
        data = [
            [split_ws.cell(row=r, column=c).value for c in range(1, 3)]
            for r in range(1, 3)
        ]
        split_wb.close()
        assert data == [["C1", "D1"], ["C2", "D2"]], \
            f"Data mismatch. Got: {data}"

    def test_split_without_label(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Split without file_label uses sheet name as filename."""
        output_dir = os.path.join(tmp_dir, "out5")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], file_label="",
        )
        assert len(created) == 1
        # Filename should be based on sheet name
        basename = os.path.basename(created[0])
        assert "Sheet1" in basename, f"Expected Sheet1 in {basename}"


# ═══════════════════════════════════════════════════════════════════════
#  2. CardSplitter.split_file — edge cases
# ═══════════════════════════════════════════════════════════════════════

class TestSplitFileEdgeCases:
    def test_non_xlsx_file(self, tmp_dir: str):
        """Non-.xlsx files should be skipped."""
        output_dir = os.path.join(tmp_dir, "out_edge1")
        txt_path = os.path.join(tmp_dir, "test.txt")
        with open(txt_path, "w") as f:
            f.write("not an xlsx")
        splitter = CardSplitter()
        created = splitter.split_file(txt_path, output_dir, ["Sheet1"], "Test")
        assert created == [], "Non-xlsx should return empty list"

    def test_sheet_not_found(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Requesting a non-existent sheet returns empty (no crash)."""
        output_dir = os.path.join(tmp_dir, "out_edge2")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["NonExistentSheet"], "Test",
        )
        assert created == [], "Non-existent sheet should return empty list"

    def test_single_sheet_file(self, tmp_dir: str, single_sheet_xlsx: str):
        """Single-sheet file should be handled (no unnecessary copy)."""
        output_dir = os.path.join(tmp_dir, "out_edge3")
        splitter = CardSplitter()
        created = splitter.split_file(
            single_sheet_xlsx, output_dir, ["OnlySheet"], "Test",
        )
        # Single sheet file — splitting should still work
        assert len(created) == 1, f"Expected 1 file, got {len(created)}"
        assert _count_xlsx_sheets(created[0]) == 1

    def test_duplicate_filename_handling(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Duplicate output filenames get a counter suffix."""
        output_dir = os.path.join(tmp_dir, "out_edge4")
        splitter = CardSplitter()
        # Split same sheet twice
        created1 = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], "TestCard",
        )
        created2 = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], "TestCard",
        )
        assert len(created1) == 1
        assert len(created2) == 1
        # Second file should have different filename (counter suffix)
        assert created1[0] != created2[0], \
            "Duplicate should produce different filename"
        assert os.path.exists(created2[0])

    def test_empty_sheet_list(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Empty sheet list returns empty."""
        output_dir = os.path.join(tmp_dir, "out_edge5")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, [], "Test",
        )
        assert created == [], "Empty sheet list should return empty"

    def test_output_dir_created(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Output directory is created if it doesn't exist."""
        output_dir = os.path.join(tmp_dir, "new_dir", "nested")
        assert not os.path.exists(output_dir)
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], "Test",
        )
        assert len(created) == 1
        assert os.path.exists(output_dir)


# ═══════════════════════════════════════════════════════════════════════
#  3. CardSplitter._extract_sheet_via_zip — content integrity
# ═══════════════════════════════════════════════════════════════════════

class TestExtractSheetViaZip:
    def test_sheet_data_correct(self, tmp_dir: str):
        """Verify extracted sheet has correct data."""
        # Create file with specific data
        sheets = {
            "DataSheet": [["Header1", "Header2"], ["Val1", "Val2"], ["Val3", "Val4"]],
            "OtherSheet": [["Other1", "Other2"]],
        }
        path = _create_multi_sheet_xlsx(tmp_dir, "data_test.xlsx", sheets)
        output_dir = os.path.join(tmp_dir, "extract1")
        os.makedirs(output_dir)

        splitter = CardSplitter()
        created = splitter.split_file(path, output_dir, ["DataSheet"], "")[0]

        # Check data
        assert _get_xlsx_cell(created, "DataSheet", 1, 1) == "Header1"
        assert _get_xlsx_cell(created, "DataSheet", 2, 2) == "Val2"
        assert _get_xlsx_cell(created, "DataSheet", 3, 1) == "Val3"

    def test_file_valid_xlsx(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Split file should be a valid .xlsx readable by openpyxl."""
        output_dir = os.path.join(tmp_dir, "extract2")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir, ["Sheet1"], "TestCard",
        )
        # openpyxl should be able to open it without errors
        wb = openpyxl.load_workbook(created[0])
        assert wb.active is not None
        wb.close()

    def test_sheet_name_preserved(self, tmp_dir: str):
        """Sheet name should be preserved in the output file."""
        sheets = {
            "СпециальноеИмя": [["A", "B"]],
            "Other": [["C", "D"]],
        }
        path = _create_multi_sheet_xlsx(tmp_dir, "name_test.xlsx", sheets)
        output_dir = os.path.join(tmp_dir, "extract3")
        os.makedirs(output_dir)

        splitter = CardSplitter()
        created = splitter.split_file(path, output_dir, ["СпециальноеИмя"], "Test")[0]

        names = openpyxl.load_workbook(created).sheetnames
        assert names == ["СпециальноеИмя"], f"Expected preserved name, got {names}"

    def test_multiple_extracts_same_source(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Extracting multiple sheets from the same source should all work."""
        output_dir = os.path.join(tmp_dir, "extract4")
        splitter = CardSplitter()
        created = splitter.split_file(
            multi_sheet_xlsx, output_dir,
            ["Sheet1", "Sheet2", "Sheet3"], "Test",
        )
        assert len(created) == 3
        for fp in created:
            assert os.path.getsize(fp) > 0, f"File {fp} is empty"


# ═══════════════════════════════════════════════════════════════════════
#  4. CardSplitter.split_many_parallel
# ═══════════════════════════════════════════════════════════════════════

class TestSplitManyParallel:
    def test_parallel_split(self, tmp_dir: str):
        """Parallel split of multiple files."""
        # Create 3 multi-sheet files
        files = []
        for i in range(3):
            sheets = {
                f"Op{i}A": [["Part", "Qty"], [f"P{i}01", "1"]],
                f"Op{i}B": [["Part", "Qty"], [f"P{i}02", "2"]],
            }
            path = _create_multi_sheet_xlsx(
                tmp_dir, f"card_{i}.xlsx", sheets,
            )
            files.append(path)

        output_dir = os.path.join(tmp_dir, "parallel_out")
        tasks = [
            (files[0], output_dir, ["Op0A", "Op0B"], "Card0"),
            (files[1], output_dir, ["Op1A", "Op1B"], "Card1"),
            (files[2], output_dir, ["Op2A", "Op2B"], "Card2"),
        ]

        splitter = CardSplitter(max_workers=2)
        created = splitter.split_many_parallel(tasks)
        assert len(created) == 6, f"Expected 6 files from 3 cards × 2 ops, got {len(created)}"
        for fp in created:
            assert os.path.exists(fp), f"File {fp} missing"
            assert _count_xlsx_sheets(fp) == 1

    def test_parallel_single_file(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Parallel split with single file should still work."""
        output_dir = os.path.join(tmp_dir, "parallel_single")
        tasks = [(multi_sheet_xlsx, output_dir, ["Sheet1"], "Card")]

        splitter = CardSplitter(max_workers=2)
        created = splitter.split_many_parallel(tasks)
        assert len(created) == 1

    def test_parallel_empty_tasks(self, tmp_dir: str):
        """Empty tasks list returns empty."""
        splitter = CardSplitter()
        created = splitter.split_many_parallel([])
        assert created == []


# ═══════════════════════════════════════════════════════════════════════
#  5. _split_file_worker
# ═══════════════════════════════════════════════════════════════════════

class TestSplitFileWorker:
    def test_worker_basic(self, tmp_dir: str, multi_sheet_xlsx: str):
        """Worker function produces correct output."""
        output_dir = os.path.join(tmp_dir, "worker_out")
        created = _split_file_worker(
            multi_sheet_xlsx, output_dir, ["Sheet1", "Sheet2"], "TestCard",
        )
        assert len(created) == 2, f"Expected 2 files, got {len(created)}"
        for fp in created:
            assert os.path.exists(fp)
            assert _count_xlsx_sheets(fp) == 1


# ═══════════════════════════════════════════════════════════════════════
#  6. Integration: split + re-parse
# ═══════════════════════════════════════════════════════════════════════

class TestSplitIntegration:
    def test_split_then_parse_detects_table(self, tmp_dir: str):
        """After splitting, each single-sheet file should still have find_part_table work."""
        # Create a card-like .xlsx with header and data
        sheets = {
            "Операция1": [
                ["物料编码", "零件名称", "数量"],
                ["P001", "Part1", "2"],
                ["P002", "Part2", "1"],
            ],
            "Операция2": [
                ["物料编码", "零件名称", "数量"],
                ["Q001", "Part3", "3"],
            ],
        }
        path = _create_multi_sheet_xlsx(tmp_dir, "card_integration.xlsx", sheets)
        output_dir = os.path.join(tmp_dir, "integ_out")

        splitter = CardSplitter()
        created = splitter.split_file(path, output_dir, ["Операция1"], "Card001")

        assert len(created) == 1
        split_path = created[0]

        # parse the split file to confirm it works
        wb = openpyxl.load_workbook(split_path)
        ws = wb.active
        result = HeuristicAnalyzer.find_part_table(ws)
        wb.close()

        assert result is not None, "find_part_table should still work on split file"
        hr, pn, qty, name = result
        assert pn == 1, f"Expected part_no=C1, got C{pn}"
        assert qty == 3, f"Expected qty=C3, got C{qty}"


# ═══════════════════════════════════════════════════════════════════════
#  7. _clean_named_ranges (XML-level unit tests)
# ═══════════════════════════════════════════════════════════════════════

class TestCleanNamedRanges:
    """Test _clean_named_ranges at the XML ElementTree level."""

    def _make_wb_root(self, defined_names: Optional[List[Dict[str, str]]] = None):
        """Create a minimal workbook.xml with definedNames."""
        import xml.etree.ElementTree as ET
        NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

        NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        root = ET.fromstring(
            f'<workbook xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
            f'  <sheets>'
            f'    <sheet name="Sheet1" sheetId="1" r:id="rId1"/>'
            f'    <sheet name="Sheet2" sheetId="2" r:id="rId2"/>'
            f'  </sheets>'
            f'</workbook>'
        )

        if defined_names:
            dn_elem = ET.SubElement(root, f'{{{NS_MAIN}}}definedNames')
            for dn in defined_names:
                d = ET.SubElement(dn_elem, f'{{{NS_MAIN}}}definedName')
                d.set('name', dn.get('name', ''))
                if 'localSheetId' in dn:
                    d.set('localSheetId', dn['localSheetId'])
                d.text = dn.get('formula', '')

        return root

    def test_removes_named_range_for_deleted_sheet(self):
        """Named range referencing deleted sheet is removed."""
        root = self._make_wb_root([
            {'name': 'MyRange', 'formula': "Sheet2!$A$1:$B$2"},
        ])
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")

        # definedNames should be empty (only range was for Sheet2)
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is None or len(dn_elem) == 0, \
            "Named range for deleted sheet should be removed"

    def test_keeps_named_range_for_kept_sheet(self):
        """Named range referencing kept sheet stays."""
        root = self._make_wb_root([
            {'name': 'MyRange', 'formula': "Sheet1!$A$1:$B$2"},
        ])
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")

        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is not None
        assert len(dn_elem) == 1, "Named range for kept sheet should remain"
        assert dn_elem[0].get('name') == 'MyRange'

    def test_removes_quoted_sheet_name(self):
        """Named range with quoted sheet name (spaces) is removed."""
        root = self._make_wb_root([
            {'name': 'Range1', 'formula': "'Sheet Two'!$A$1"},
        ])
        _clean_named_ranges(root, {"Sheet Two"}, "Sheet1")
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is None or len(dn_elem) == 0

    def test_updates_local_sheet_id(self):
        """Non-zero localSheetId on kept sheet's named ranges is reset to 0."""
        root = self._make_wb_root([
            {'name': 'LocalRange', 'formula': "Sheet1!$A$1",
             'localSheetId': '1'},
        ])
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is not None
        assert dn_elem[0].get('localSheetId') == '0', \
            "localSheetId should be reset to 0"

    def test_removes_elem_when_empty(self):
        """definedNames element removed entirely if no ranges remain."""
        root = self._make_wb_root([
            {'name': 'ToDelete', 'formula': "Sheet2!$A$1"},
        ])
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is None, "Empty definedNames should be removed"

    def test_no_defined_names_at_all(self):
        """Workbook without definedNames is unchanged."""
        root = self._make_wb_root()  # no defined names
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is None

    def test_mixed_ranges_keeps_and_deletes(self):
        """Mixed: keep ranges for remaining sheet, delete for removed sheets."""
        root = self._make_wb_root([
            {'name': 'Keep', 'formula': "Sheet1!$A$1"},
            {'name': 'Delete', 'formula': "Sheet2!$B$2"},
            {'name': 'AlsoKeep', 'formula': "Sheet1!$C$3"},
        ])
        _clean_named_ranges(root, {"Sheet2"}, "Sheet1")
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        dn_elem = root.find('m:definedNames', ns)
        assert dn_elem is not None
        names = [d.get('name') for d in dn_elem]
        assert "Keep" in names
        assert "AlsoKeep" in names
        assert "Delete" not in names
