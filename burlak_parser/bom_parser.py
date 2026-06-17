"""Модуль чтения BOM-файла (Bill of Materials / Ведомость материалов).

Формат: .xlsx (таблица на китайском/английском/русском языках).

Алгоритм работы:
  1. Загружает .xlsx и обходит ВСЕ листы.
  2. Для каждого листа использует эвристический анализатор для поиска:
     - Строки заголовков
     - Колонок с парт-номерами, названиями и количествами
     - Колонок комплектаций
  3. Строит ГЛОБАЛЬНЫЙ словарь парт-номеров и названий (сканирует ВСЕ строки,
     а не только для конкретной комплектации).
  4. Извлекает количества по каждой комплектации.
  5. Агрегирует данные по всем листам.

Универсален — не привязан к конкретным моделям автомобилей, брендам или
форматам. Использует эвристический анализатор из heuristic_analyzer.py.

Класс BOMService — обёртка для использования в FastAPI/серверной архитектуре.
"""

from __future__ import annotations

import logging
import os
import tempfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Set, Tuple

import openpyxl

from burlak_parser.heuristic_analyzer import (
    HeuristicAnalyzer,
    clean_cell_text,
)
from burlak_parser.normalizer import (
    normalize_quantity,
    normalize_part_number,
    clean_part_number,
    is_valid_part_number,
)

# Листы с мета-именами, которые НЕ создают конфигурации (单车用量, 发动机附件 и т.д.)
_NON_CONFIG_SHEET_KEYWORDS = (
    "单车用量", "组件数量", "发动机附件",
    "Расход на один автомобиль", "количество компонентов",
)

logger = logging.getLogger(__name__)


@dataclass
class PartInfo:
    """Информация о детали из BOM."""
    part_number: str
    name_cn: str = ""
    name_en: str = ""
    # Количество для конкретной комплектации (будет заполнено после выбора)
    quantity: float = 0.0
    # Номера/коды комплектаций, для которых указана деталь
    applicable_configs: List[str] = field(default_factory=list)


@dataclass
class BOMData:
    """Результат парсинга BOM-файла."""
    parts: Dict[str, PartInfo]  # part_number -> PartInfo
    config_names: List[str]  # названия колонок комплектаций
    config_quantities: Dict[str, Dict[str, float]]  # config_name -> {part_number -> qty}
    source_file: str = ""
    # Глобальный словарь названий (составлен из ВСЕХ строк, а не только для комплектации)
    global_names: Dict[str, Tuple[str, str]] = field(default_factory=dict)  # part_number -> (name_cn, name_en)


def parse_bom(file_path: str) -> BOMData:
    """Разобрать BOM-файл и вернуть структурированные данные.

    Args:
        file_path: Путь к .xlsx файлу BOM.

    Returns:
        BOMData со всеми извлечёнными данными.
    """
    logger.info("Загрузка BOM-файла: %s", file_path)

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        logger.warning("Не удалось загрузить обычным режимом (%s), пробуем read_only", e)
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

    try:
        sheet_names = wb.sheetnames

        # ── Результаты, агрегированные по всем листам ──
        all_parts: Dict[str, PartInfo] = {}
        all_config_quantities: Dict[str, Dict[str, float]] = {}
        all_config_names: List[str] = []
        all_global_names: Dict[str, Tuple[str, str]] = {}
        seen_config_names: Dict[str, str] = {}  # config_name -> нормализованный оригинал

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            logger.info(
                "Анализ листа: %s (строк: %s, колонок: %s)",
                sheet_name, ws.max_row, ws.max_column,
            )

            # Проверяем, является ли лист BOM-кандидатом
            is_bom = HeuristicAnalyzer.is_sheet_bom_candidate(
                ws, min_configs=2, sheet_name=sheet_name,
            )
            if not is_bom:
                logger.info("Лист не является BOM-кандидатом, пропуск: %s", sheet_name)
                continue

            # ── 1. Поиск строки заголовков ──
            header_rows = HeuristicAnalyzer.find_header_rows(ws)
            if not header_rows:
                logger.warning("Не найдена строка заголовков в листе: %s", sheet_name)
                continue

            # ── 2. Определение типов колонок ──
            col_types = HeuristicAnalyzer.detect_column_types(ws, header_rows)
            part_no_col = col_types.get("part_no", 0)
            name_cn_col = col_types.get("name_cn", 0)
            name_en_col = col_types.get("name_en", 0)

            if part_no_col == 0:
                logger.warning("Не найдена колонка парт-номеров в листе: %s", sheet_name)
                continue

            header_row = header_rows[0]

            # ── 3. Строим ГЛОБАЛЬНЫЙ словарь названий (ВСЕ строки, ВСЕ листы) ──
            sheet_names_dict = HeuristicAnalyzer.build_global_name_dict(
                ws, part_no_col, name_cn_col, name_en_col, header_row,
            )
            for pn, (nc, ne) in sheet_names_dict.items():
                if pn not in all_global_names:
                    all_global_names[pn] = (nc, ne)
                else:
                    existing_cn, existing_en = all_global_names[pn]
                    if not existing_cn and nc:
                        existing_cn = nc
                    if not existing_en and ne:
                        existing_en = ne
                    all_global_names[pn] = (existing_cn, existing_en)

            # ── 4. Определяем колонки комплектаций ──
            config_cols = HeuristicAnalyzer.detect_config_columns(ws, header_rows, col_types)
            qty_col = col_types.get("qty", 0)

            # ── 5. Если есть отдельная qty-колонка (спец-листы 附件) ──
            is_non_config_sheet = any(kw in sheet_name for kw in _NON_CONFIG_SHEET_KEYWORDS)

            if (not config_cols or len(config_cols) == 0) and qty_col > 0 and not is_non_config_sheet:
                data_start = header_row + 1
                config_name = sheet_name
                seen_config_names[config_name] = config_name
                all_config_names.append(config_name)
                all_config_quantities[config_name] = {}

                for row_idx in range(data_start, (ws.max_row or data_start) + 1):
                    pn = HeuristicAnalyzer.get_cell_value(ws, row_idx, part_no_col)
                    if pn is None:
                        continue
                    pn_str = clean_cell_text(pn)
                    if not pn_str or pn_str.startswith("~$"):
                        continue
                    if not is_valid_part_number(pn_str):
                        continue

                    qty_val = HeuristicAnalyzer.get_cell_value(ws, row_idx, qty_col)
                    qty = normalize_quantity(qty_val)

                    if qty > 0:
                        pn_normalized = clean_part_number(pn_str)
                        current_qty = all_config_quantities[config_name].get(pn_normalized, 0.0)
                        all_config_quantities[config_name][pn_normalized] = current_qty + qty

                        if pn_normalized not in all_parts:
                            all_parts[pn_normalized] = PartInfo(part_number=pn_str)
                        if config_name not in all_parts[pn_normalized].applicable_configs:
                            all_parts[pn_normalized].applicable_configs.append(config_name)

                logger.info(
                    "Лист %s: спец-лист с qty-колонкой, %d деталей",
                    sheet_name, len(all_config_quantities[config_name]),
                )
                continue

            # ── 5b. Non-config sheets (单车用量, 发动机附件) — collect parts only ──
            if is_non_config_sheet and qty_col > 0:
                data_start = header_row + 1
                for row_idx in range(data_start, (ws.max_row or data_start) + 1):
                    pn = HeuristicAnalyzer.get_cell_value(ws, row_idx, part_no_col)
                    if pn is None:
                        continue
                    pn_str = clean_cell_text(pn)
                    if not pn_str or pn_str.startswith("~$"):
                        continue
                    if not is_valid_part_number(pn_str):
                        continue
                    pn_normalized = clean_part_number(pn_str)
                    if pn_normalized not in all_parts:
                        all_parts[pn_normalized] = PartInfo(part_number=pn_str)
                logger.info("Лист %s: не-конфигурационный, детали собраны в all_parts", sheet_name)
                continue

            if not config_cols:
                logger.info("Лист %s: не найдено колонок комплектаций, пропуск", sheet_name)
                continue

            # ── 6. Дедупликация имён комплектаций ──
            config_names: List[str] = []
            for col_idx in config_cols:
                name = HeuristicAnalyzer.get_cell_value(ws, header_row, col_idx)
                name_str = str(name) if name is not None else ""
                name_str = name_str.replace("\n", " ").replace("\r", "").strip()

                if not name_str:
                    for look_row in range(max(1, header_row - 1), 0, -1):
                        meta_val = HeuristicAnalyzer.get_cell_value(ws, look_row, col_idx)
                        if meta_val is not None:
                            meta_str = str(meta_val).strip()
                            if meta_str and len(meta_str) < 80:
                                name_str = meta_str
                                break
                if not name_str:
                    name_str = f"Config_{col_idx}"
                config_names.append(name_str)

            deduped_indices: List[int] = []
            seen_norm: Set[str] = set()
            for i, name in enumerate(config_names):
                norm = name.lower().replace(" ", "").replace("-", "")
                if norm not in seen_norm:
                    seen_norm.add(norm)
                    deduped_indices.append(i)

            if len(deduped_indices) < len(config_cols):
                logger.info(
                    "Дедупликация: %d -> %d имён комплектаций",
                    len(config_cols), len(deduped_indices),
                )
                config_cols = [config_cols[i] for i in deduped_indices]
                config_names = [config_names[i] for i in deduped_indices]

            # ── 7. Парсинг данных комплектаций ──
            data_start = header_row + 1
            max_row = ws.max_row or data_start
            sheet_config_count = 0

            for row_idx in range(data_start, max_row + 1):
                pn = HeuristicAnalyzer.get_cell_value(ws, row_idx, part_no_col)
                if pn is None:
                    continue
                pn_str = clean_cell_text(pn)
                if not pn_str or pn_str.startswith("~$"):
                    continue
                if not is_valid_part_number(pn_str):
                    continue

                pn_normalized = clean_part_number(pn_str)

                if pn_normalized not in all_parts:
                    all_parts[pn_normalized] = PartInfo(part_number=pn_str)

                part = all_parts[pn_normalized]

                for i, col_idx in enumerate(config_cols):
                    config_val = str(HeuristicAnalyzer.get_cell_value(ws, row_idx, col_idx) or '').strip()

                    if config_val.upper() == 'S' and qty_col > 0:
                        qty = normalize_quantity(HeuristicAnalyzer.get_cell_value(ws, row_idx, qty_col))
                    elif config_val in ('-', '–', '—', ''):
                        continue
                    else:
                        qty = normalize_quantity(config_val)

                    if qty > 0:
                        config_name = config_names[i]
                        if config_name not in seen_config_names:
                            seen_config_names[config_name] = config_name
                            all_config_names.append(config_name)

                        if config_name not in all_config_quantities:
                            all_config_quantities[config_name] = {}

                        current_qty = all_config_quantities[config_name].get(pn_normalized, 0.0)
                        all_config_quantities[config_name][pn_normalized] = current_qty + qty

                        if config_name not in part.applicable_configs:
                            part.applicable_configs.append(config_name)

                        sheet_config_count += 1

            logger.info(
                "Лист %s: BOM, %d колонок комплектаций, %d строк с данными",
                sheet_name, len(config_cols), sheet_config_count,
            )

    finally:
        wb.close()

    # ── Финальная агрегация ──
    logger.info("Загружено деталей (уникальных): %d", len(all_parts))
    logger.info("Найдено комплектаций: %d", len(all_config_names))
    logger.info("Глобальный словарь названий: %d записей", len(all_global_names))

    for cn in all_config_names[:5]:
        qty_count = len(all_config_quantities.get(cn, {}))
        logger.info("  %s: %d деталей", cn[:50], qty_count)
    if len(all_config_names) > 5:
        logger.info("  ... и ещё %d комплектаций", len(all_config_names) - 5)

    # Применяем глобальные названия к деталям, у которых нет названия
    for pn, part in all_parts.items():
        if (not part.name_cn and not part.name_en) and pn in all_global_names:
            gc, ge = all_global_names[pn]
            if not part.name_cn and gc:
                part.name_cn = gc
            if not part.name_en and ge:
                part.name_en = ge

    return BOMData(
        parts=all_parts,
        config_names=all_config_names,
        config_quantities=all_config_quantities,
        source_file=file_path,
        global_names=all_global_names,
    )


def get_config_quantities(bom: BOMData, config_name: str) -> Dict[str, PartInfo]:
    """Получить данные деталей для выбранной комплектации.

    Args:
        bom: Распарсенные BOM-данные.
        config_name: Название комплектации.

    Returns:
        Словарь {part_number: PartInfo} с заполненным quantity для комплектации.
    """
    if config_name not in bom.config_quantities:
        raise ValueError(
            f"Комплектация '{config_name}' не найдена. "
            f"Доступные: {bom.config_names[:10]}..."
        )

    result: Dict[str, PartInfo] = {}
    for part_no, qty in bom.config_quantities[config_name].items():
        if part_no in bom.parts:
            part = bom.parts[part_no]
            result[part_no] = PartInfo(
                part_number=part_no,
                name_cn=part.name_cn,
                name_en=part.name_en,
                quantity=qty,
            )
        else:
            # Берём из глобального словаря
            gc, ge = bom.global_names.get(part_no, ("", ""))
            result[part_no] = PartInfo(
                part_number=part_no,
                name_cn=gc,
                name_en=ge,
                quantity=qty,
            )

    return result


def get_all_config_quantities(bom: BOMData) -> Dict[str, Dict[str, PartInfo]]:
    """Получить данные деталей для ВСЕХ комплектаций одновременно.

    Args:
        bom: Распарсенные BOM-данные.

    Returns:
        Словарь {config_name: {part_number: PartInfo}}.
    """
    return {cn: get_config_quantities(bom, cn) for cn in bom.config_names}


def lookup_part_name(bom: BOMData, part_number: str) -> Tuple[str, str]:
    """Найти название детали по парт-номеру.

    Сначала ищет в parts, затем в global_names.

    Args:
        bom: BOM-данные.
        part_number: Парт-номер.

    Returns:
        (name_cn, name_en)
    """
    if part_number in bom.parts:
        part = bom.parts[part_number]
        if part.name_cn or part.name_en:
            return (part.name_cn, part.name_en)
    return bom.global_names.get(part_number, ("", ""))


class BOMService:
    """Сервис парсинга BOM-файлов.

    Готов к использованию в серверной архитектуре (FastAPI).
    Поддерживает:
      - Загрузку из файла (load)
      - Загрузку из памяти (load_from_bytes) — для HTTP upload
      - Асинхронную загрузку (load_async) — не блокирует event loop
      - Автоочистку временных файлов (cleanup / context manager)
    """

    def __init__(self):
        self._bom: Optional[BOMData] = None
        self._temp_paths: List[str] = []

    @property
    def bom(self) -> Optional[BOMData]:
        return self._bom

    @property
    def is_loaded(self) -> bool:
        return self._bom is not None

    def load(self, file_path: str) -> BOMData:
        """Загрузить и распарсить BOM-файл.

        Args:
            file_path: Путь к .xlsx файлу BOM.

        Returns:
            Распарсенные данные BOMData.
        """
        self._bom = parse_bom(file_path)
        return self._bom

    def load_from_bytes(self, data: bytes, filename: str = "bom.xlsx") -> BOMData:
        """Загрузить BOM из байтового содержимого (in-memory upload).

        Сохраняет данные во временный файл, парсит, возвращает результат.
        Временный файл будет удалён при вызове cleanup() или выходе из
        контекстного менеджера.

        Args:
            data: Байтовое содержимое .xlsx файла.
            filename: Имя файла для определения расширения.

        Returns:
            Распарсенные данные BOMData.
        """
        suffix = os.path.splitext(filename)[1] or ".xlsx"
        fd, path = tempfile.mkstemp(suffix=suffix, prefix="bom_upload_")
        os.close(fd)
        with open(path, "wb") as f:
            f.write(data)
        self._temp_paths.append(path)
        return self.load(path)

    async def load_async(self, data: bytes, filename: str = "bom.xlsx") -> BOMData:
        """Асинхронная загрузка BOM из байтов.

        Парсинг CPU-bound — выполняется в отдельном потоке,
        не блокируя event loop.

        Args:
            data: Байтовое содержимое .xlsx файла.
            filename: Имя файла для определения расширения.

        Returns:
            Распарсенные данные BOMData.
        """
        import asyncio
        return await asyncio.to_thread(self.load_from_bytes, data, filename)

    def cleanup(self) -> None:
        """Удалить все временные файлы, созданные при load_from_bytes."""
        for path in self._temp_paths:
            try:
                if os.path.isfile(path):
                    os.remove(path)
            except Exception:
                pass
        self._temp_paths.clear()

    def __enter__(self) -> BOMService:
        return self

    def __exit__(self, *args: object) -> None:
        self.cleanup()

    def get_config_names(self) -> List[str]:
        """Получить список названий всех найденных комплектаций."""
        if not self._bom:
            raise RuntimeError("BOM не загружен. Вызовите load() сначала.")
        return list(self._bom.config_names)

    def get_config_count(self) -> int:
        """Получить количество найденных комплектаций."""
        if not self._bom:
            return 0
        return len(self._bom.config_names)

    def get_parts_for_config(self, config_name: str) -> Dict[str, PartInfo]:
        """Получить детали для конкретной комплектации."""
        if not self._bom:
            raise RuntimeError("BOM не загружен. Вызовите load() сначала.")
        return get_config_quantities(self._bom, config_name)

    def get_all_configs(self) -> Dict[str, Dict[str, PartInfo]]:
        """Получить детали для ВСЕХ комплектаций."""
        if not self._bom:
            raise RuntimeError("BOM не загружен. Вызовите load() сначала.")
        return get_all_config_quantities(self._bom)

    def get_all_part_numbers(self) -> Set[str]:
        """Получить множество ВСЕХ уникальных парт-номеров из BOM."""
        if not self._bom:
            return set()
        return set(self._bom.parts.keys())

    def lookup_name(self, part_number: str) -> Tuple[str, str]:
        """Найти название детали по парт-номеру (с учётом глобального словаря)."""
        if not self._bom:
            return ("", "")
        return lookup_part_name(self._bom, part_number)
